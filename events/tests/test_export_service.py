import csv
import io
from datetime import date

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from events.models import Event, EventCategory, EventFormat, EventLocation, EventMember
from events.services.export_service import ExportService
from shop.models import Order, OrderItem


class ExportServiceTestBase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.category = EventCategory.objects.create(name="testcat")
        cls.eventformat = EventFormat.objects.create(name="testformat")
        cls.location = EventLocation.objects.create(title="testloc")

        cls.event = Event.objects.create(
            name="Test Event",
            category=cls.category,
            eventformat=cls.eventformat,
            location=cls.location,
            label="test-event",
            price="100.00",
        )

        cls.member1 = EventMember.objects.create(
            event=cls.event,
            firstname="Max",
            lastname="Mustermann",
            email="max@example.com",
            attend_status="registered",
            member_type="V",
        )

        cls.member2 = EventMember.objects.create(
            event=cls.event,
            firstname="Erika",
            lastname="Musterfrau",
            email="erika@example.com",
            attend_status="registered",
            member_type="F",
        )

        cls.member_waiting = EventMember.objects.create(
            event=cls.event,
            firstname="Hans",
            lastname="Wartend",
            email="hans@example.com",
            attend_status="waiting",
            member_type="V",
        )

        cls.service = ExportService()


class ExportMembersCsvTest(ExportServiceTestBase):

    def test_response_content_type(self):
        response = self.service.export_members_csv("test-event")
        self.assertEqual(response["Content-Type"], "text/csv")

    def test_response_filename(self):
        response = self.service.export_members_csv("test-event")
        today = date.today()
        self.assertIn(f"test-event_TN_{today}.csv", response["Content-Disposition"])

    def test_csv_header_row(self):
        response = self.service.export_members_csv("test-event")
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        header = next(reader)
        self.assertEqual(header, ["Vorname", "Nachname", "E-Mail", "Datum", "Mitgliedschaft"])

    def test_csv_contains_all_members(self):
        response = self.service.export_members_csv("test-event")
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        # header + 3 members (registered + waiting)
        self.assertEqual(len(rows), 4)

    def test_csv_member_data(self):
        response = self.service.export_members_csv("test-event")
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        next(reader)  # skip header
        first_row = next(reader)
        self.assertEqual(first_row[0], "Max")
        self.assertEqual(first_row[1], "Mustermann")
        self.assertEqual(first_row[2], "max@example.com")


class ExportMvMembersCsvTest(ExportServiceTestBase):

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.member1.vote_transfer = "Delegiert an Test"
        cls.member1.save()

    def test_response_content_type(self):
        response = self.service.export_mv_members_csv("test-event")
        self.assertEqual(response["Content-Type"], "text/csv")

    def test_filter_by_firstname(self):
        response = self.service.export_mv_members_csv(
            "test-event", filters={"firstname": "Max"}
        )
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        # header + 1 matching member
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][0], "Max")

    def test_filter_by_lastname(self):
        response = self.service.export_mv_members_csv(
            "test-event", filters={"lastname": "Musterfrau"}
        )
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][1], "Musterfrau")

    def test_filter_by_email(self):
        response = self.service.export_mv_members_csv(
            "test-event", filters={"email": "erika"}
        )
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        self.assertEqual(len(rows), 2)

    def test_filter_vote_transfer_yes(self):
        response = self.service.export_mv_members_csv(
            "test-event", filters={"vote_transfer_yes": True}
        )
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        # header + 1 member with vote_transfer set
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][0], "Max")

    def test_filter_vote_transfer_no(self):
        response = self.service.export_mv_members_csv(
            "test-event", filters={"vote_transfer_no": True}
        )
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        # header + 2 members without vote_transfer
        self.assertEqual(len(rows), 3)

    def test_csv_header_without_vote_transfer(self):
        response = self.service.export_mv_members_csv("test-event")
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        header = next(reader)
        self.assertEqual(
            header,
            ["Vorname", "Nachname", "E-Mail", "Status", "Datum", "Mitgliedschaft"],
        )


class ExportMoodleParticipantsTest(ExportServiceTestBase):

    def test_response_content_type(self):
        response = self.service.export_moodle_participants(self.event)
        self.assertEqual(response["Content-Type"], "text/csv")

    def test_csv_header(self):
        response = self.service.export_moodle_participants(self.event)
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        header = next(reader)
        self.assertEqual(
            header,
            ["username", "firstname", "lastname", "email", "course1", "role1"],
        )

    def test_only_registered_members(self):
        response = self.service.export_moodle_participants(self.event)
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        # header + 2 registered members (not the waiting one)
        self.assertEqual(len(rows), 3)

    def test_moodle_row_format(self):
        response = self.service.export_moodle_participants(self.event)
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        next(reader)  # skip header
        row = next(reader)
        # username = email, then firstname, lastname, email, course1=label, role1=student
        self.assertEqual(row[0], row[3])  # username == email
        self.assertEqual(row[4], "test-event")
        self.assertEqual(row[5], "student")

    def test_filename_contains_label_and_date(self):
        response = self.service.export_moodle_participants(self.event)
        self.assertIn("test-event", response["Content-Disposition"])
        self.assertIn(str(date.today()), response["Content-Disposition"])


class ExportParticipantsXlsTest(ExportServiceTestBase):

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.order1 = Order.objects.create(
            firstname="Max",
            lastname="Mustermann",
            email="max@example.com",
            payment_type="r",
            payment_date=timezone.now(),
        )
        OrderItem.objects.create(
            order=cls.order1,
            event=cls.event,
            price=100,
            premium_price=135,
            cost=100,
        )
        cls.order2 = Order.objects.create(
            firstname="Erika",
            lastname="Musterfrau",
            email="erika@example.com",
            payment_type="r",
            payment_date=timezone.now(),
        )
        OrderItem.objects.create(
            order=cls.order2,
            event=cls.event,
            price=100,
            premium_price=135,
            cost=100,
        )

    def test_participants_version_content_type(self):
        response = self.service.export_participants_xls(self.event, "participants")
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_participants_version_filename(self):
        response = self.service.export_participants_xls(self.event, "participants")
        self.assertIn("TeilnehmerInnen_test-event", response["Content-Disposition"])

    def test_controlling_version_filename(self):
        response = self.service.export_participants_xls(self.event, "controlling")
        self.assertIn("Controlling_test-event", response["Content-Disposition"])

    def test_participants_version_has_xlsx_content(self):
        response = self.service.export_participants_xls(self.event, "participants")
        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws.title, "TeilnehmerInnen")
        # header row + 2 registered members
        self.assertEqual(ws.max_row, 3)

    def test_controlling_version_has_xlsx_content(self):
        response = self.service.export_participants_xls(self.event, "controlling")
        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws.title, "Controlling")
        # Should have registered members + waitlist section + speakers + event details
        self.assertTrue(ws.max_row > 3)
